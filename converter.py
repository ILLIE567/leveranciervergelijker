"""Lokale bestandslezing en vaste controles, los van de gebruikersinterface."""

import csv
import re
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import pandas as pd
from openpyxl import load_workbook

FIELDS = ("Title", "Variant SKU", "Variant Price", "Variant Inventory Qty")
ERROR_COLUMNS = ["Rijnummer", "Type fout", "Uitleg"]
MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 50_000
MAX_COLUMNS = 100
ALIASES = {
    "Title": ["title", "productnaam", "naam", "product name", "name", "omschrijving"],
    "Variant SKU": ["variant sku", "sku", "artikelnummer", "productcode", "article number", "code"],
    "Variant Price": ["variant price", "prijs", "price", "verkoopprijs"],
    "Variant Inventory Qty": ["variant inventory qty", "voorraad", "stock", "quantity", "aantal"],
}


def check_headers(headers):
    if not headers or any(not isinstance(h, str) or not h.strip() for h in headers):
        raise ValueError("Elke kolom moet een niet-lege tekstnaam hebben op de eerste rij.")
    normalized = [h.strip().casefold() for h in headers]
    if len(normalized) != len(set(normalized)):
        raise ValueError("Dubbele kolomnamen gevonden. Maak de namen eerst uniek.")
    if len(headers) > MAX_COLUMNS:
        raise ValueError(f"Maximaal {MAX_COLUMNS} kolommen toegestaan.")


def read_supplier_file(data, filename, delimiter=None, encoding="utf-8-sig", sheet_name=None):
    """Lees bytes in geheugen; bewaar CSV-tekst exact en voer Excel-formules niet uit."""
    if not data:
        raise ValueError("Het bestand is leeg.")
    if len(data) > MAX_BYTES:
        raise ValueError("Het bestand is te groot. Maximaal 10 MB toegestaan.")
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        try:
            text = data.decode(encoding)
        except (UnicodeDecodeError, LookupError) as exc:
            raise ValueError("Tekencodering niet leesbaar. Kies UTF-8 of Windows-1252.") from exc
        if "\x00" in text:
            raise ValueError("Dit bestand lijkt geen tekst-CSV te zijn.")
        if delimiter is None:
            try:
                delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",;\t").delimiter
            except csv.Error as exc:
                raise ValueError("Scheidingsteken niet herkend. Kies het handmatig.") from exc
        if delimiter not in (",", ";", "\t"):
            raise ValueError("Kies komma, puntkomma of tab als scheidingsteken.")
        rows = []
        try:
            reader = csv.reader(StringIO(text, newline=""), delimiter=delimiter, strict=True)
            headers = next(reader)
            check_headers(headers)
            for record_number, row in enumerate(reader, start=2):
                # Een lege CSV-regel blijft een datarij en krijgt later fouten.
                if not row:
                    row = [""] * len(headers)
                if len(row) != len(headers):
                    raise ValueError(f"CSV-record {record_number}: {len(row)} velden, verwacht {len(headers)}. Controleer scheidingstekens en quotes.")
                if len(rows) >= MAX_ROWS:
                    raise ValueError(f"Maximaal {MAX_ROWS} datarijen toegestaan; er wordt niets afgekapt.")
                rows.append(row)
        except (csv.Error, StopIteration) as exc:
            raise ValueError("Ongeldige of lege CSV. Controleer de aanhalingstekens en kopregel.") from exc
        return pd.DataFrame(rows, columns=headers, dtype=object)
    if extension == ".xlsx":
        try:
            with ZipFile(BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError("Excel-bestand is uitgepakt te groot (maximaal 50 MB).")
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        except (BadZipFile, OSError, KeyError) as exc:
            raise ValueError("Dit bestand is geen leesbaar .xlsx-bestand.") from exc
        try:
            if len(workbook.sheetnames) > 1 and sheet_name is None:
                raise ValueError("Meerdere werkbladen gevonden. Kies eerst een werkblad.")
            if sheet_name is not None and sheet_name not in workbook.sheetnames:
                raise ValueError("Het gekozen werkblad bestaat niet.")
            sheet = workbook[sheet_name or workbook.sheetnames[0]]
            # Sommige leveranciers schrijven onjuiste of ontbrekende dimensies.
            # Lees de echte cellen; vertrouw niet op het opgegeven werkbladbereik.
            sheet.reset_dimensions()
            iterator = sheet.iter_rows()
            header_cells = next(iterator, ())
            headers = [cell.value for cell in header_cells]
            check_headers(headers)
            rows = []
            for cells in iterator:
                if len(rows) >= MAX_ROWS:
                    raise ValueError(f"Maximaal {MAX_ROWS} datarijen toegestaan; er wordt niets afgekapt.")
                if len(cells) > len(headers):
                    raise ValueError(f"Excel-rij {len(rows) + 2} bevat cellen buiten de kolomnamen. Geef elke kolom een naam.")
                values = []
                for cell in cells:
                    if cell.data_type == "f":
                        raise ValueError(f"Excel-formule op {cell.coordinate}. Vervang formules door vaste waarden in een kopie.")
                    if cell.data_type == "e":
                        raise ValueError(f"Excel-foutwaarde op {cell.coordinate}. Corrigeer deze cel in een kopie.")
                    value = "" if cell.value is None else cell.value
                    # We bewaren het type: numerieke Excel-SKU's worden later geblokkeerd.
                    values.append(value)
                rows.append(values + [""] * (len(headers) - len(values)))
            return pd.DataFrame(rows, columns=headers, dtype=object)
        finally:
            workbook.close()
    raise ValueError("Alleen .csv en .xlsx worden ondersteund.")


def excel_sheet_names(data):
    if len(data) > MAX_BYTES:
        raise ValueError("Het bestand is te groot. Maximaal 10 MB toegestaan.")
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("Excel-bestand is uitgepakt te groot (maximaal 50 MB).")
        workbook = load_workbook(BytesIO(data), read_only=True, keep_links=False)
        try:
            return workbook.sheetnames
        finally:
            workbook.close()
    except (BadZipFile, OSError, KeyError) as exc:
        raise ValueError("Dit bestand is geen leesbaar .xlsx-bestand.") from exc


def suggest_mapping(columns):
    suggestions = {}
    for field, aliases in ALIASES.items():
        matches = [column for column in columns if column.strip().casefold() in aliases]
        suggestions[field] = matches[0] if len(matches) == 1 else None
    return suggestions


def is_empty(value):
    return value is None or (isinstance(value, str) and not value.strip()) or bool(pd.isna(value))


def parse_number(value, kind):
    """Geen duizendtallen, valuta, exponenten of stil afronden."""
    if is_empty(value):
        raise ValueError(f"{kind} ontbreekt.")
    if isinstance(value, bool):
        raise ValueError(f"{kind} moet een getal zijn, geen ja/nee-waarde.")
    text = str(value).strip()
    if not re.fullmatch(r"-?[0-9]+(?:[.,][0-9]+)?", text):
        raise ValueError(f"{kind} is geen geldig getal. Gebruik bijvoorbeeld 12,95 of 12.95, zonder valuta of duizendtalscheiding.")
    number = Decimal(text.replace(",", "."))
    if number < 0:
        raise ValueError(f"{kind} mag niet negatief zijn.")
    if kind == "Prijs":
        if number.as_tuple().exponent < -2:
            raise ValueError("Prijs heeft meer dan twee decimalen; de app rondt niet automatisch af.")
        return "0.00" if number == 0 else f"{number:.2f}"
    if number.as_tuple().exponent < -2:
        raise ValueError("Voorraad heeft meer dan twee decimalen of een mogelijke duizendtalscheiding. Gebruik een geheel getal zonder scheiding.")
    if number != number.to_integral_value():
        raise ValueError("Voorraad moet een geheel getal zijn.")
    return str(int(number))


def convert_products(frame, mapping):
    """Behoud iedere rij. Bij fouten is de uitvoer alleen een controlepreview."""
    if any(not mapping.get(field) or mapping[field] not in frame.columns for field in FIELDS):
        raise ValueError("Koppel alle vier Shopify-velden aan bestaande leverancierskolommen.")
    if len(set(mapping[field] for field in FIELDS)) != len(FIELDS):
        raise ValueError("Gebruik voor elk Shopify-veld een andere leverancierskolom.")
    errors = []
    products = []
    seen_skus = {}

    def error(row_number, error_type, explanation):
        errors.append([row_number, error_type, explanation])

    for row_number, (_, row) in enumerate(frame.iterrows(), start=2):
        product = {}
        for field in ("Title", "Variant SKU"):
            value = row[mapping[field]]
            product[field] = "" if is_empty(value) else str(value)
            label = "SKU" if field == "Variant SKU" else "Productnaam"
            if is_empty(value):
                error(row_number, f"Ontbrekende {label.lower()}", f"{label} ontbreekt.")
            elif not isinstance(value, str):
                error(row_number, f"Ongeldige {label.lower()}", f"{label} is geen tekst. Sla deze Excel-cel op als tekst. Voorloopnullen kunnen al verloren zijn gegaan; de app verzint ze niet.")
            elif value != value.strip():
                error(row_number, f"Spaties bij {label.lower()}", f"{label} heeft spaties aan het begin of einde. Pas dit aan in een kopie van het bronbestand.")
            elif any(ord(char) < 32 or ord(char) == 127 for char in value):
                error(row_number, f"Controlekarakter in {label.lower()}", f"{label} bevat een tab, regeleinde of ander controlekarakter.")
            elif value.startswith(("=", "+", "-", "@")):
                error(row_number, f"Onveilige {label.lower()}", f"{label} begint met een spreadsheet-formuleteken. Controleer dit handmatig; de app wijzigt de tekst niet.")
        sku = product["Variant SKU"]
        if sku.strip():
            seen_skus.setdefault(sku.strip().casefold(), []).append(row_number)
        for field, kind in (("Variant Price", "Prijs"), ("Variant Inventory Qty", "Voorraad")):
            try:
                product[field] = parse_number(row[mapping[field]], kind)
            except (ValueError, ArithmeticError) as exc:
                product[field] = ""
                error(row_number, f"Ongeldige {kind.lower()}", str(exc) if isinstance(exc, ValueError) else f"{kind} valt buiten het ondersteunde getalbereik.")
        products.append(product)
    for row_numbers in seen_skus.values():
        if len(row_numbers) > 1:
            for row_number in row_numbers:
                error(row_number, "Dubbele SKU", "SKU komt meerdere keren voor (ook hoofdletterverschillen/spaties tellen mee). Controleer alle betrokken rijen.")
    output = pd.DataFrame(products, columns=FIELDS, dtype=object)
    # Vaste technische velden voor losse producten, nooit gegokte brongegevens.
    output["Handle"] = ["supplier-" + sku.encode("utf-8").hex() for sku in output["Variant SKU"]]
    output["Option1 Name"] = "Title"
    output["Option1 Value"] = "Default Title"
    output["Variant Inventory Tracker"] = "shopify"
    output["Variant Inventory Policy"] = "deny"
    output["Variant Fulfillment Service"] = "manual"
    report = pd.DataFrame(errors, columns=ERROR_COLUMNS)
    if not report.empty:
        report = report.sort_values("Rijnummer", kind="stable").reset_index(drop=True)
    return output, report


def csv_bytes(frame):
    return frame.to_csv(index=False, lineterminator="\n").encode("utf-8-sig")


def export_shopify(output, report):
    if not report.empty:
        raise ValueError("Download geblokkeerd: los eerst alle fouten op.")
    if output.empty:
        raise ValueError("Geen producten om te downloaden.")
    return csv_bytes(output)
